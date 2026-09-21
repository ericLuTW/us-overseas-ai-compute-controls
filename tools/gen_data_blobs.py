#!/usr/bin/env python3
"""Regenerate the dashboard's embedded data blobs (window.__DATA_ZH / __DATA_EN)
from the source workbooks, and splice them into index.html.

The dashboard (index.html) is a self-contained Claude-Design React app whose
data lives in two inline <script> blocks:
    <script>window.__DATA_ZH = {...}</script>
    <script>window.__DATA_EN = {...}</script>

This script rebuilds those JSON objects from the ZH/EN workbooks. Three
structures are NOT derivable from the workbooks and are carried over verbatim
from the current index.html (they were authored at design time):
    GLOSS   - 66-entry abbreviation glossary (predates the v2 framework 詞彙表)
    nodes   - N1..N15 supply-chain node display names + group assignment
    GRP_EN  - node-group EN label map

Usage:
    python3 tools/gen_data_blobs.py check   <zh.xlsx> <en.xlsx>   # regression: compare against current index.html blobs
    python3 tools/gen_data_blobs.py splice  <zh.xlsx> <en.xlsx>   # rewrite index.html in place
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
INDEX = ROOT / "index.html"

# per-language column headers: blob field -> header name
POL_COLS = {
    "zh": {"code": "policy_id", "item": "條目（中文）", "id": "編號", "date": "日期",
           "prog": "進度", "who": "頒布者／提案人", "law": "法源", "bg": "背景與脈絡",
           "body": "內容", "links": "來源連結", "disp": "新碼"},
    "en": {"code": "policy_id", "item": "Entry (English)", "id": "Citation", "date": "Date",
           "prog": "Status", "who": "Issuer / Sponsor", "law": "Legal Basis",
           "bg": "Background & Context", "body": "Content", "links": "Source Links",
           "disp": "New Code"},
}
TOOL_COLS = {
    "zh": {"tool": "政策工具", "act": "工具作用", "appl": "適用對象", "dir": "作用方向",
           "macro": "工具大類", "grp": "工具群組", "trig": "觸發要件／門檻",
           "op": "操作方式", "spec": "具體標的", "law": "授權／制定主體",
           "impl": "實施機關", "note": "備註"},
    "en": {"tool": "Policy Tool", "act": "Tool Function", "appl": "Applicable Parties",
           "dir": "Direction", "macro": "Tool Category", "grp": "Tool Group",
           "trig": "Trigger Conditions / Thresholds", "op": "Mode of Operation",
           "spec": "Specific Targets", "law": "Authorizing / Issuing Body",
           "impl": "Implementing Agency", "note": "Notes"},
}
SHEETS = {
    "zh": {"pol": "政策條目", "tools": "政策工具盤點", "cat": "標的×工具長表",
           "node": "節點×工具長表", "polnode": "政策×節點長表", "vocab": "詞彙表"},
    "en": {"pol": "Policy Entries", "tools": "Policy Tool Inventory", "cat": "Targets × Tools",
           "node": "Nodes × Tools", "polnode": "Policies × Nodes", "vocab": "Glossary"},
}


def sheet_rows(ws):
    hdr = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    out = []
    for r in range(2, ws.max_row + 1):
        row = {h: ws.cell(row=r, column=c).value for h, c in hdr.items() if h}
        if any(v not in (None, "") for v in row.values()):
            out.append(row)
    return out


def s(v):
    if v is None:
        return ""
    return str(v).strip()


def extract_current_blobs(html):
    blobs = {}
    for lang in ("ZH", "EN"):
        m = re.search(r"<script>window\.__DATA_" + lang + r" = ", html)
        start = m.end()
        end = html.index("</script>", start)
        blobs[lang.lower()] = json.loads(html[start:end].rstrip().rstrip(";"))
    return blobs


def build(path, lang, statics):
    wb = openpyxl.load_workbook(path, data_only=True)
    sh = SHEETS[lang]
    pc, tc = POL_COLS[lang], TOOL_COLS[lang]

    pol_rows = sheet_rows(wb[sh["pol"]])
    tool_rows = sheet_rows(wb[sh["tools"]])
    cat_rows = sheet_rows(wb[sh["cat"]])
    node_rows = sheet_rows(wb[sh["node"]])
    vocab_rows = sheet_rows(wb[sh["vocab"]])

    # policies
    pol, links, code2, track_of = [], {}, {}, {}
    pol_by_code = {}
    for r in pol_rows:
        code = s(r["policy_id"])
        disp = s(r[pc["disp"]])
        date = s(r[pc["date"]])
        p = {"code": code, "disp": disp, "item": s(r[pc["item"]]), "id": s(r[pc["id"]]),
             "year": date[:4], "date": date, "prog": s(r[pc["prog"]]),
             "who": s(r[pc["who"]]), "law": s(r[pc["law"]]), "bg": s(r[pc["bg"]]),
             "body": s(r[pc["body"]]), "track": disp.split("-")[0] if disp else ""}
        pol.append(p)
        pol_by_code[code] = p
        if s(r[pc["links"]]):
            # workbook separates multiple URLs with ' ｜ '; the blob uses newlines
            links[code] = re.sub(r"\s*｜\s*", "\n", s(r[pc["links"]]))
        code2[code] = disp
        track_of[code] = p["track"]

    # tools (row order defines numeric tool no 1..N)
    tools = {}
    code_of = {}
    for i, r in enumerate(tool_rows, 1):
        pid = s(r["policy_id"])
        parent = pol_by_code[pid]
        t = {"no": i, "tool": s(r[tc["tool"]]), "dir": s(r[tc["dir"]]),
             "macro": s(r[tc["macro"]]), "grp": s(r[tc["grp"]]),
             "prog": parent["prog"], "item": parent["item"],
             "spec": s(r[tc["spec"]]), "act": s(r[tc["act"]]), "op": s(r[tc["op"]]),
             "law": s(r[tc["law"]]), "appl": s(r[tc["appl"]]), "trig": s(r[tc["trig"]]),
             "impl": s(r[tc["impl"]]), "note": s(r[tc["note"]])}
        tools[s(r["tool_id"])] = t
        code_of[str(i)] = pid

    def light(t):
        return {"no": t["no"], "tool": t["tool"], "dir": t["dir"]}

    # cells: target_category ||| macro -> [full tool]
    cells = {}
    tool_cats = {}
    for r in cat_rows:
        t = tools[s(r["tool_id"])]
        cat = s(r["target_category"])
        cells.setdefault(cat + "|||" + t["macro"], []).append(t)
        tool_cats.setdefault(t["no"], []).append(cat)

    # nodecells: node ||| macro -> [light]; and tool -> nodes
    # node_id "0" marks tools without a supply-chain node -> excluded everywhere
    nodecells = {}
    tool_nodes = {}
    for r in node_rows:
        nid = s(r["node_id"])
        if nid == "0":
            continue
        t = tools[s(r["tool_id"])]
        nodecells.setdefault(nid + "|||" + t["macro"], []).append(light(t))
        tool_nodes.setdefault(t["no"], []).append(nid)

    # polcells: policy ||| node -> [light] (derived from node×tool via tool's policy)
    polcells = {}
    for r in node_rows:
        nid = s(r["node_id"])
        if nid == "0":
            continue
        t = tools[s(r["tool_id"])]
        pid = code_of[str(t["no"])]
        polcells.setdefault(pid + "|||" + nid, []).append(light(t))

    # bridge: target_category ||| node -> [light] (tools mapped to both)
    bridge = {}
    seen_bridge = set()
    for no, cats in tool_cats.items():
        for cat in cats:
            for nid in tool_nodes.get(no, []):
                key = cat + "|||" + nid
                if (key, no) in seen_bridge:
                    continue
                seen_bridge.add((key, no))
                t = next(t for t in tools.values() if t["no"] == no)
                bridge.setdefault(key, []).append(light(t))

    # vocab-driven lists
    def vocab(domain):
        return [r for r in vocab_rows if s(r["vocab_domain"]) == domain]

    name_key = "name_zh" if lang == "zh" else "name_en"
    macros = [s(r[name_key]) for r in vocab("tool_major")]
    cats = [s(r[name_key]) for r in vocab("target_category")]
    macros_en = [s(r["name_en"]) for r in vocab("tool_major")]
    cats_en = [s(r["name_en"]) for r in vocab("target_category")]
    tracks = [{"code": s(r["code"]), "zh": s(r["name_zh"]), "en": s(r["name_en"])}
              for r in vocab("track")]

    trk_idx = {tr["code"]: i for i, tr in enumerate(tracks)}
    pol.sort(key=lambda p: (trk_idx.get(p["track"], 99), p["disp"]))

    return {
        "D": {"macros": macros, "cats": cats, "nodes": statics["nodes"],
              "cells": cells, "nodecells": nodecells, "polcells": polcells,
              "bridge": bridge, "pol": pol},
        "CODE_OF": code_of, "LINKS": links, "GRP_EN": statics["GRP_EN"],
        "GLOSS": statics["GLOSS"], "trackOf": track_of, "TRACKS": tracks,
        "cats_en": cats_en, "macros_en": macros_en, "CODE2": code2,
    }


def deep_diff(a, b, path="", out=None, limit=40):
    if out is None:
        out = []
    if len(out) >= limit:
        return out
    if type(a) is not type(b):
        out.append(f"{path}: type {type(a).__name__} != {type(b).__name__}")
    elif isinstance(a, dict):
        for k in a.keys() | b.keys():
            if k not in a:
                out.append(f"{path}.{k}: extra in generated")
            elif k not in b:
                out.append(f"{path}.{k}: missing in generated")
            else:
                deep_diff(a[k], b[k], f"{path}.{k}", out, limit)
    elif isinstance(a, list):
        if len(a) != len(b):
            out.append(f"{path}: len {len(a)} != {len(b)}")
        for i, (x, y) in enumerate(zip(a, b)):
            deep_diff(x, y, f"{path}[{i}]", out, limit)
    elif a != b:
        out.append(f"{path}: {str(a)[:60]!r} != {str(b)[:60]!r}")
    return out


def main():
    mode, zh_path, en_path = sys.argv[1], sys.argv[2], sys.argv[3]
    html = INDEX.read_text(encoding="utf-8")
    current = extract_current_blobs(html)
    statics = {k: current["zh"][k] for k in ("GLOSS", "GRP_EN")}
    statics["nodes"] = current["zh"]["D"]["nodes"]

    gen = {"zh": build(zh_path, "zh", statics), "en": build(en_path, "en", statics)}

    if mode == "check":
        ok = True
        for lang in ("zh", "en"):
            diffs = deep_diff(current[lang], gen[lang], lang)
            print(f"[{lang}] diffs: {len(diffs)}")
            for d in diffs:
                print("   ", d)
            ok = ok and not diffs
        sys.exit(0 if ok else 1)

    if mode == "splice":
        for lang in ("ZH", "EN"):
            blob = json.dumps(gen[lang.lower()], ensure_ascii=False,
                              separators=(",", ":")).replace("</", "<\\/")
            m = re.search(r"<script>window\.__DATA_" + lang + r" = ", html)
            start = m.end()
            end = html.index("</script>", start)
            html = html[:start] + blob + html[end:]
        INDEX.write_text(html, encoding="utf-8")
        print("spliced index.html")
        for lang in ("zh", "en"):
            d = gen[lang]
            print(f"  {lang}: pol={len(d['D']['pol'])} tools={len(d['CODE_OF'])} "
                  f"cells={len(d['D']['cells'])} tracks={[t['code'] for t in d['TRACKS']]}")
        sys.exit(0)

    print("unknown mode")
    sys.exit(2)


if __name__ == "__main__":
    main()
